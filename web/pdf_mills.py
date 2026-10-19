"""Parse a mill-list PDF into company farms with 32-sided polygons.

UML IDs are resolved through the Universal Mill List workbook (UML ID, Latitude,
Longitude). A point location becomes a regular 32-gon so the map can treat the
mill like a property outline. Analysis is not started here.
"""
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

MILL_RADIUS_M = 10000
MILL_SIDES = 32

_HEADER_KEYS = (
    ("uml_id", ("uml id", "uml")),
    ("lat", ("latitude", "lat")),
    ("lon", ("longitude", "lon", "lng", "long")),
    ("gps", ("gps coordinates", "gps", "coordinates")),
    ("name", ("mill name",)),
    ("mill_company", ("mill company", "parent company")),
    ("supplier", ("direct supplier", "supplier")),
    ("country", ("country",)),
)
_UML_CELL = re.compile(r"PO\d{4,}", re.IGNORECASE)
_PAIR = re.compile(r"(-?\d{1,3}\.\d+)\s*[,;/]\s*(-?\d{1,3}\.\d+)")
_COMPANY = re.compile(
    r"([A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9&.'’\-]{0,40}(?:[ \t]+[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9&.'’\-]{0,40}){0,5})['’]s supply chain",
    re.IGNORECASE,
)


def circle_polygon(lon, lat, radius_m=MILL_RADIUS_M, sides=MILL_SIDES):
    """Regular polygon around a point, in degrees, closed GeoJSON ring."""
    cos_lat = math.cos(math.radians(lat))
    if abs(cos_lat) < 0.01:
        cos_lat = 0.01 if cos_lat >= 0 else -0.01
    dlat = radius_m / 111_320.0
    dlon = radius_m / (111_320.0 * cos_lat)
    ring = []
    for index in range(sides):
        angle = 2 * math.pi * index / sides
        ring.append([
            round(lon + dlon * math.cos(angle), 6),
            round(lat + dlat * math.sin(angle), 6),
        ])
    ring.append(ring[0])
    return {"type": "Polygon", "coordinates": [ring]}


def ring_bbox(ring):
    lons = [pt[0] for pt in ring]
    lats = [pt[1] for pt in ring]
    return [min(lons), min(lats), max(lons), max(lats)]


def _number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _lat_lon(lat, lon):
    lat = _number(lat)
    lon = _number(lon)
    if lat is None or lon is None:
        return None
    if abs(lat) > 90 and abs(lon) <= 90:
        lat, lon = lon, lat
    if abs(lat) > 90 or abs(lon) > 180:
        return None
    return lat, lon


def _pair(text):
    if not text:
        return None
    match = _PAIR.search(str(text))
    if not match:
        return None
    return _lat_lon(match.group(1), match.group(2))


def _header_key(text):
    cleaned = re.sub(r"\s+", " ", (text or "").strip().lower())
    for key, aliases in _HEADER_KEYS:
        if cleaned in aliases:
            return key
    return None


def _items(page):
    found = []

    def visitor(text, _cm, tm, _font, _size):
        cleaned = (text or "").replace("\u00a0", " ").strip()
        if cleaned:
            found.append({"x": float(tm[4]), "y": float(tm[5]), "text": cleaned})

    page.extract_text(visitor_text=visitor)
    return found


def _rows(items, y_tol=2.0):
    ordered = sorted(items, key=lambda item: (-item["y"], item["x"]))
    rows = []
    for item in ordered:
        if rows and abs(rows[-1]["y"] - item["y"]) <= y_tol:
            rows[-1]["cells"].append(item)
        else:
            rows.append({"y": item["y"], "cells": [item]})
    for row in rows:
        row["cells"].sort(key=lambda cell: cell["x"])
    return rows


def _headers_from(cells):
    headers = []
    for cell in cells:
        key = _header_key(cell["text"])
        if key:
            headers.append({"key": key, "x": cell["x"]})
    keys = {header["key"] for header in headers}
    if "uml_id" not in keys and "lat" not in keys and "gps" not in keys:
        return None
    if len(headers) < 2:
        return None
    return headers


def _assign(cells, headers):
    fields = {}
    for cell in cells:
        header = min(headers, key=lambda item: abs(item["x"] - cell["x"]))
        fields.setdefault(header["key"], [])
        fields[header["key"]].append(cell["text"])
    return {key: " ".join(parts).strip() for key, parts in fields.items()}


def _uml_id(fields, cells):
    candidate = fields.get("uml_id") or ""
    match = _UML_CELL.fullmatch(candidate.strip())
    if match:
        return match.group(0).upper()
    for cell in cells:
        match = _UML_CELL.fullmatch(cell["text"].strip())
        if match:
            return match.group(0).upper()
    return None


def _coordinates(fields):
    pair = _lat_lon(fields.get("lat"), fields.get("lon"))
    if pair:
        return pair
    return _pair(fields.get("gps"))


def load_uml(path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        header = [re.sub(r"\s+", " ", str(cell or "").strip().lower()) for cell in next(iterator)]

        def column(*names):
            for name in names:
                if name in header:
                    return header.index(name)
            return None

        id_col = column("uml id", "uml_id", "id")
        lat_col = column("latitude", "lat")
        lon_col = column("longitude", "lon", "lng")
        if id_col is None or lat_col is None or lon_col is None:
            raise ValueError("UML workbook needs UML ID, Latitude, and Longitude columns.")
        index = {}
        for row in iterator:
            if not row or id_col >= len(row):
                continue
            match = _UML_CELL.search(str(row[id_col] or ""))
            if not match:
                continue
            pair = _lat_lon(
                row[lat_col] if lat_col < len(row) else None,
                row[lon_col] if lon_col < len(row) else None,
            )
            if pair:
                index[match.group(0).upper()] = pair
        return index
    finally:
        workbook.close()


def _company_name(text, filename):
    match = _COMPANY.search(text or "")
    if match:
        words = re.sub(r"\s+", " ", match.group(1)).strip(" -–—").split()
        while len(words) > 1 and any(ch.isdigit() for ch in words[0]):
            words.pop(0)
        if words:
            return " ".join(words)
    stem = Path(filename or "company").stem
    stem = re.sub(r"[-_]+", " ", stem).strip()
    return stem[:80] or "Company"


def _farm_id(uml_id, name, lat, lon, index):
    if uml_id:
        return uml_id.lower()
    slug = unicodedata.normalize("NFKD", name or "mill")
    slug = "".join(ch for ch in slug if not unicodedata.combining(ch))
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", slug).strip("-").lower()[:40] or "mill"
    if lat is not None and lon is not None:
        return f"{slug}-{lat:.4f}-{lon:.4f}"
    return f"{slug}-{index}"


def parse_pdf(pdf_path, uml_path, filename="upload.pdf"):
    from pypdf import PdfReader

    uml_index = load_uml(uml_path)
    reader = PdfReader(str(pdf_path))
    headers = None
    texts = []
    farms = []
    seen = set()
    for page in reader.pages:
        items = _items(page)
        texts.extend(item["text"] for item in items)
        for row in _rows(items):
            found = _headers_from(row["cells"])
            if found:
                headers = found
                continue
            if not headers:
                continue
            fields = _assign(row["cells"], headers)
            uml_id = _uml_id(fields, row["cells"])
            coords = _coordinates(fields)
            source = "pdf" if coords else None
            if coords is None and uml_id and uml_id in uml_index:
                coords = uml_index[uml_id]
                source = "uml"
            if not uml_id and not coords:
                continue
            name = fields.get("name") or fields.get("mill_company") or uml_id or "Mill"
            lat = coords[0] if coords else None
            lon = coords[1] if coords else None
            farm_id = _farm_id(uml_id, name, lat, lon, len(farms) + 1)
            if farm_id in seen:
                continue
            seen.add(farm_id)
            geometry = None
            bbox = None
            if coords:
                geometry = circle_polygon(lon, lat)
                bbox = ring_bbox(geometry["coordinates"][0])
            farms.append({
                "id": farm_id,
                "name": name,
                "mill_company": fields.get("mill_company") or "",
                "supplier": fields.get("supplier") or "",
                "country": fields.get("country") or "",
                "uml_id": uml_id,
                "lat": lat,
                "lon": lon,
                "coordinate_source": source,
                "resolved": coords is not None,
                "radius_m": MILL_RADIUS_M if coords else None,
                "sides": MILL_SIDES if coords else None,
                "geometry": geometry,
                "bbox": bbox,
            })
    if not texts:
        raise ValueError("No text could be read from that PDF.")
    if not farms:
        raise ValueError("No UML IDs or coordinates were found in that PDF.")
    resolved = sum(1 for farm in farms if farm["resolved"])
    return {
        "company": _company_name("\n".join(texts), filename),
        "farms": farms,
        "mill_count": len(farms),
        "resolved_count": resolved,
        "unresolved_count": len(farms) - resolved,
    }


def main(argv):
    if len(argv) < 4 or argv[1] != "parse":
        print("usage: pdf_mills.py parse PDF UML_XLSX FILENAME", file=sys.stderr)
        return 2
    try:
        payload = parse_pdf(Path(argv[2]), Path(argv[3]), argv[4] if len(argv) > 4 else Path(argv[2]).name)
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1
    json.dump(payload, sys.stdout)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
